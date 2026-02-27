#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
初始化 data.xlsx 的 Sheet0（index 配置）
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

DATA_XLSX = Path(__file__).resolve().parents[1] / 'res' / 'data.xlsx'

# 18 个 section 的基础配置
SECTIONS_CONFIG = [
    {
        '主标题': '世界奇妙物语系列',
        '概要': '包含TV版三季及1990-2025年特别篇，涵盖日本最经典的悬疑短篇集锦剧',
        '封面图片路径': './res/covers/01.webp',
        '卡片数量': 82
    },
    {
        '主标题': '藤原龙也系列',
        '概要': '日本实力派演员藤原龙也主演作品集，包含《死亡笔记》《大逃杀》等经典',
        '封面图片路径': './res/covers/02.webp',
        '卡片数量': 39
    },
    {
        '主标题': '憨豆先生系列',
        '概要': '英国喜剧大师罗温·艾金森经典作品，无需语言也能笑到肚子疼',
        '封面图片路径': './res/covers/03.webp',
        '卡片数量': 6
    },
    {
        '主标题': '生化危机系列',
        '概要': '米拉·乔沃维奇主演的动作科幻大片，改编自同名游戏的经典电影系列',
        '封面图片路径': './res/covers/04.webp',
        '卡片数量': 10
    },
    {
        '主标题': '哈利波特系列',
        '概要': 'J.K.罗琳魔法世界的完整电影系列，陪伴一代人成长的魔法传奇',
        '封面图片路径': './res/covers/05.webp',
        '卡片数量': 8
    },
    {
        '主标题': '死神来了系列',
        '概要': '死神要你三更死，不得留人到五更！经典惊悚系列，环环相扣的死亡链条，笑到喷血浆的离奇死亡方式',
        '封面图片路径': './res/covers/06.webp',
        '卡片数量': 6
    },
    {
        '主标题': '数码暴龙系列',
        '概要': '暑假被电脑拐跑到怪兽服务器，小学生一脚踏进会蓝屏的虚拟世界，输出全靠喊，要如何拯救两个世界？',
        '封面图片路径': './res/covers/07.webp',
        '卡片数量': 28
    },
    {
        '主标题': '推理探案系列',
        '概要': '死神小学生几十年如一日在凶案现场破案，名侦探柯南和金田一带你走进柯学的杂技世界',
        '封面图片路径': './res/covers/08.webp',
        '卡片数量': 42
    },
    {
        '主标题': '异世界系列',
        '概要': '今天死在异世界的是哪一回？带着只能读档不能存档的废柴能力，体验一场用命拼出来的压力试炼',
        '封面图片路径': './res/covers/09.webp',
        '卡片数量': 17
    },
    {
        '主标题': '竞速赛车系列',
        '概要': '深夜送豆腐顺便虐车神，山路发卡弯当直线开，藤原拓海带你体验一场靠重力和排水渠极限下坡的死亡公路之旅',
        '封面图片路径': './res/covers/10.webp',
        '卡片数量': 14
    },
    {
        '主标题': '超人英雄系列',
        '概要': '认真就会输的秃头英雄，每天打工顺手把世界清零，埼玉带你见识无敌也无法避免面临财务危机的英雄社会学',
        '封面图片路径': './res/covers/11.webp',
        '卡片数量': 9
    },
    {
        '主标题': '游戏王系列',
        '概要': '我100点生命稳如老狗、你10000点生命犹如风中残烛，口胡打牌也能拯救世界？代代相传的不仅是名字和发型，还有印刷绝技！',
        '封面图片路径': './res/covers/12.webp',
        '卡片数量': 37
    },
    {
        '主标题': '三大民工漫',
        '概要': '打工人下班追了十年热血，死神火影海贼王陪你见证从拼修行到拼爹开挂的热血进化史',
        '封面图片路径': './res/covers/13.webp',
        '卡片数量': 26
    },
    {
        '主标题': 'JOJO 系列',
        '概要': '一条家族血脉打八代怪人，从英国鬼故事一路打到宇宙歌剧，JOJO 把"气氛到位就算合理"演成视觉圣经',
        '封面图片路径': './res/covers/14.webp',
        '卡片数量': 0
    },
    {
        '主标题': '宠物小精灵系列',
        '概要': '十岁就被赶出家环游世界，把小动物塞进精灵球里互相斗殴，离家出走一走就是 25 年终于走到功成名就',
        '封面图片路径': './res/covers/15.webp',
        '卡片数量': 0
    },
    {
        '主标题': '哆啦 A 梦系列',
        '概要': '来自 22 世纪的蓝胖子专门用道具帮小学生逃避人生，"如果有这种道具就好了"成为一代人的童年妄想实录',
        '封面图片路径': './res/covers/16.webp',
        '卡片数量': 0
    },
    {
        '主标题': '精选动画系列',
        '概要': '90 后必看动漫合集！仙境传说、Re0、Fate、钢炼、死亡笔记、东京喰种等热血治愈神作全收录',
        '封面图片路径': './res/covers/98.webp',
        '卡片数量': 29
    },
    {
        '主标题': '其他系列',
        '概要': '不容错过的独立佳作！《打机王》怀旧经典、《疯狂动物城》温馨治愈，每部都是精品',
        '封面图片路径': './res/covers/99.webp',
        '卡片数量': 15
    }
]

def init_data_xlsx():
    """初始化 data.xlsx"""
    headers = ['主标题', '概要', '封面图片路径', '卡片数量']
    
    # 检查文件是否存在
    if DATA_XLSX.exists():
        print(f"📖 加载现有 Excel: {DATA_XLSX}")
        wb = load_workbook(DATA_XLSX)
        # 删除 Sheet0 如果存在
        if 'Sheet0' in wb.sheetnames:
            del wb['Sheet0']
    else:
        print(f"📝 创建新 Excel 文件")
        wb = Workbook()
        # 移除默认 sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

    # 创建 Sheet0
    ws = wb.create_sheet('Sheet0', 0)
    
    # 添加表头
    ws.append(headers)
    header_row = ws[1]
    for cell in header_row:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 添加数据
    for config in SECTIONS_CONFIG:
        row = [config.get(h, '') for h in headers]
        ws.append(row)
    
    # 调整列宽
    ws.column_dimensions['A'].width = 20  # 主标题
    ws.column_dimensions['B'].width = 40  # 概要
    ws.column_dimensions['C'].width = 30  # 封面图片路径
    ws.column_dimensions['D'].width = 12  # 卡片数量
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    DATA_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DATA_XLSX)
    
    print(f"✅ 已初始化 Sheet0")
    print(f"📁 文件: {DATA_XLSX}")
    print(f"📊 包含 {len(SECTIONS_CONFIG)} 个 section 配置")


if __name__ == '__main__':
    init_data_xlsx()
