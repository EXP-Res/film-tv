#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
更新 data.xlsx 结构，添加标签列表列
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DATA_FILE = './res/data.xlsx'

# 标签数据（对应各个 section）
TAGS_DATA = {
    '世界奇妙物语系列': ['悬疑', '惊悚', 'SP特别篇'],
    '藤原龙也系列': ['高分经典', '实力派演员'],
    '憨豆先生系列': ['喜剧经典', '全球爆笑'],
    '生化危机系列': ['动作科幻', '游戏改编'],
    '哈利波特系列': ['魔法奇幻', '小说改编'],
    '死神来了系列': ['惊悚悬疑', '连环死亡'],
    '数码暴龙系列': ['童年经典', '数码进化'],
    '推理探案系列': ['柯学神话', '动作推理'],
    '异世界系列': ['死亡回归', '后宫养成'],
    '竞速赛车系列': ['死亡竞速', '山路飘移'],
    '超人英雄系列': ['无敌神话', '反套路英雄'],
    '游戏王系列': ['决斗王道', '嘴炮抽卡'],
    '三大民工漫': ['热血王道', '血统开挂'],
    'JOJO 系列': ['中二名场面', '妖异美学'],
    '宠物小精灵系列': ['精灵对战', '收服养成'],
    '哆啦 A 梦系列': ['梦幻道具', '搞笑友情'],
    '精选短篇动画': ['热血神作', '青春回忆'],
    '精选影视': ['独立精品', '高分推荐'],
}

def main():
    print("更新 data.xlsx...")
    wb = load_workbook(DATA_FILE)
    ws = wb['index']
    
    # 添加标签列头（第 5 列）
    ws.cell(1, 5, '标签列表')
    header_cell = ws.cell(1, 5)
    header_cell.font = Font(bold=True, color='FFFFFF')
    header_cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 为每一行添加标签数据
    for row_idx in range(2, ws.max_row + 1):
        title = ws.cell(row_idx, 1).value
        if title and title in TAGS_DATA:
            tags = ','.join(TAGS_DATA[title])
            ws.cell(row_idx, 5, tags)
    
    # 调整列宽
    ws.column_dimensions['E'].width = 20
    
    wb.save(DATA_FILE)
    print(f"✅ 已更新 {DATA_FILE}")
    print(f"✓ 添加了标签列表列")

if __name__ == '__main__':
    main()
