#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 Excel 生成 HTML 页面脚本
功能：
1. 从 Sheet0 生成 index.html
2. 从 Sheet1-n 生成 section-01.html ~ section-n.html
"""

import os
import re
import random
import string
from openpyxl import load_workbook
from jinja2 import Template
from color_log.clog import log

# ===========================
# 配置
# ===========================
CHASET = 'utf-8'
SPLIT_LINE = "=" * 60
SHEET_INDEX = 'index'
EXCEL_FILE = "./res/data.xlsx"  # Excel 文件名
OUTPUT_DIR = "sections"  # 输出目录
TEMPLATE_SECTION = "./templates/section-tpl.html"  # section 模板文件
TEMPLATE_CARD = "./templates/card-tpl.html"  # card 模板文件
TEMPLATE_INDEX = "./templates/index-tpl.html"  # index 模板文件
TEMPLATE_COVER = "./templates/cover-tpl.html"  # cover 模板文件
OUTPUT_INDEX = "index.html"  # index 输出文件


# ===========================
# 主流程
# ===========================
def main():
    log.info(SPLIT_LINE)
    log.info("HTML 生成器 - 从 Excel 生成网站页面")
    log.info(SPLIT_LINE)
    
    # 检查 Excel 文件
    if not os.path.exists(EXCEL_FILE):
        log.info(f"❌ 错误: 找不到 {EXCEL_FILE}")
        log.info(f"请在项目根目录放置 {EXCEL_FILE} 文件")
        return
    
    # 加载 Excel 数据
    log.info(f"📖 读取 Excel 文件: {EXCEL_FILE}")
    sheets = load_excel_data(EXCEL_FILE)
    log.info(f"✓ 找到 {len(sheets)} 个 Sheet")
    
    # 创建输出目录
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    log.info(f"✓ 输出目录: {OUTPUT_DIR}/")
    
    # 处理每个 Sheet
    sheet_list = list(sheets.keys())
    index_sheet = SHEET_INDEX if SHEET_INDEX in sheets else (sheet_list[0] if sheet_list else None)
    section_sheets = [s for s in sheet_list if s != SHEET_INDEX]
    
    generate_sections(sheets, section_sheets)
    if index_sheet:
        generate_index(sheets, index_sheet)
    
    # 更新 JS 文件配置
    update_js_config(sheets, section_sheets)
    
    log.info(SPLIT_LINE)
    log.info("✅ 生成完成！")
    log.info(SPLIT_LINE)
    
    
        

# Sheet 1-n 生成 section HTML
def generate_sections(sheets, sheet_list) :
    for idx, sheet_name in enumerate(sheet_list):
        idx += 1
        html_name = f'section-{idx:02d}.html'
        log.info(f"📄 开始转换 Sheet {sheet_name} -> {html_name}")
        
        is_ok, header, lines = load_sheet_table(sheets, sheet_name)
        if not is_ok:
            log.warn(f"⚠️  {sheet_name} 没有数据，跳过")
            continue
        # log.debug(header)
        # log.debug(lines)
        
        
        # 生成 HTML
        html_content = generate_section(lines, sheet_name, idx)
        
        # 保存文件
        html_path = os.path.join(OUTPUT_DIR, html_name)
        with open(html_path, 'w', encoding=CHASET) as f:
            f.write(html_content)
        
        log.info(f"✓ 生成 {len(lines)} 个卡片")
        log.info(f"✓ 已保存到: {html_path}")


def generate_section(sheet_data, section_title, section_number):
    """生成 section HTML（从模板渲染）"""
    # 生成卡片 HTML
    cards = []
    for data in sheet_data:
        card = generate_card_html(data)
        cards.append(card)
        # log.debug(f"\n{card}")
    cards_html = '\n\n'.join(cards)
    
    # 准备模板上下文
    context = {
        'section_title': section_title,
        'section_number': section_number,
        'cards': cards_html
    }
    
    # 加载 section 模板并渲染
    tpl_text = load_template_section()
    tpl = Template(tpl_text)
    return tpl.render(**context)


def generate_card_html(card_data):
    """生成单个电影卡片 HTML（从模板渲染）"""

    # 准备字段
    context = {
        'poster_path': card_data.get('封面图片路径', '../res/placeholder.webp'), 
        'title': card_data.get('主标题', ''),
        'subtitle': card_data.get('副标题', ''),
        'description': card_data.get('概要', ''),
        'language': card_data.get('语言', '日语'),
        'subtitle_text': card_data.get('字幕', '中文'),
        'menu_path': card_data.get('目录路径', ''), 
        'download_name': card_data.get('网盘名称', '百度网盘'),
        'download_link': card_data.get('下载链接', 'FIXME'),
        'password': card_data.get('解压密码', 'FIXME'),
        'formats': card_data.get('支持格式', 'mp4'),
        'card_foot': card_data.get('卡片页脚', '')
    }

    # 生成格式徽章（仍保留原有解析逻辑）
    context['format_badges'] = parse_formats(context['formats'])

    # 加载 card 模板并渲染
    tpl_text = load_template_card()
    tpl = Template(tpl_text)
    return tpl.render(**context)



# Sheet 0 生成 index.html 
def generate_index(sheets, sheet_name) :
    log.info(f"📄 开始转换 Sheet 0 -> index.html")

    is_ok, header, lines = load_sheet_table(sheets, sheet_name)
    if not is_ok:
        log.warn(f"⚠️  {sheet_name} 没有数据，跳过")
        return
    
    # 生成 cover 卡片 HTML
    covers = []
    for idx, data in enumerate(lines, 1):
        cover = generate_cover_html(data, idx)
        covers.append(cover)
    covers_html = '\n\n'.join(covers)
    
    # 准备模板上下文
    context = {
        'covers': covers_html
    }
    
    # 加载 index 模板并渲染
    tpl_text = load_template_index()
    tpl = Template(tpl_text)
    html_content = tpl.render(**context)
    
    # 保存文件
    with open(OUTPUT_INDEX, 'w', encoding=CHASET) as f:
        f.write(html_content)
    
    log.info(f"✓ 生成 {len(lines)} 个 section 卡片")
    log.info(f"✓ 已保存到: {OUTPUT_INDEX}")


def generate_cover_html(cover_data, idx):
    """生成 section cover 卡片 HTML（从模板渲染）"""
    # 准备字段
    tags_str = cover_data.get('标签列表', '')
    tags = re.split(r'[,，;；]', tags_str)
    tags = [f'<i class="bi bi-star"></i> {t.strip()}<br/>' for t in tags if t.strip()]
    tags_list = '\n                                  '.join(tags)
    
    context = {
        'section_id': idx,
        'section_link': f"./sections/section-{idx:02d}.html",
        'poster_path': cover_data.get('封面图片路径', f'./res/covers/{idx:02d}.webp'),
        'section_name': cover_data.get('主标题', ''),
        'description': cover_data.get('概要', ''),
        'count': cover_data.get('卡片数量', 0),
        'tags': tags_list,
    }
    
    # 加载 cover 模板并渲染
    tpl_text = load_template_cover()
    tpl = Template(tpl_text)
    return tpl.render(**context)



def load_sheet_table(sheets, sheet_name) :
    table = sheets[sheet_name]
    header = table['headers']
    lines = table['data']
    if not lines:
        log.warn(f"⚠️  {sheet_name} 没有数据，跳过")
        return False, [], []
    # log.debug(header)
    # log.debug(lines)
    return (True, header, lines)



# ===========================
# 工具函数
# ===========================
def generate_uuid(length=8):
    """生成随机 UUID (8 字符)"""
    return ''.join(random.choices(string.ascii_letters + string.digits, k=length))


def get_badge_html(format_name):
    """生成格式徽章 HTML"""
    badge_colors = {
        'mp4': ('blue', 'mp4'),
        'mkv': ('purple', 'mkv'),
        'rmvb': ('orange', 'rmvb'),
        'tv': ('green', 'tv'),
        'avi': ('red', 'avi'),
        'flv': ('yellow', 'flv'),
    }
    
    color, label = badge_colors.get(format_name.lower(), ('gray', format_name))
    return f'<img src="https://img.shields.io/badge/{label}-Yes-{color}.svg" data-bs-toggle="tooltip" title="包含 {label} 视频格式">'


def parse_formats(format_str):
    """解析格式列表，返回 HTML"""
    if not format_str or format_str.strip() == '':
        return ''
    
    formats = re.split(r'[,，;；]', format_str)
    formats = [f.strip() for f in formats if f.strip()]
    return '\n                      '.join([get_badge_html(f) for f in formats])


def load_excel_data(excel_file):
    """加载 Excel 数据"""
    if not os.path.exists(excel_file):
        raise FileNotFoundError(f"找不到文件: {excel_file}")
    
    wb = load_workbook(excel_file)
    sheets = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = []
        
        # 获取标题行
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        # 读取数据行
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            row_data = {}
            for col_idx, cell in enumerate(row):
                header = headers[col_idx]
                row_data[header] = cell.value if cell.value is not None else ''
            data.append(row_data)
        
        sheets[sheet_name] = {
            'headers': headers,
            'data': data
        }
    
    return sheets


def load_template_section():
    """加载 section 模板"""
    with open(TEMPLATE_SECTION, 'r', encoding=CHASET) as f:
        return f.read()


def load_template_card():
    """加载 card 模板"""
    with open(TEMPLATE_CARD, 'r', encoding=CHASET) as f:
        return f.read()


def load_template_index():
    """加载 index 模板"""
    with open(TEMPLATE_INDEX, 'r', encoding=CHASET) as f:
        return f.read()


def load_template_cover():
    """加载 cover 模板"""
    with open(TEMPLATE_COVER, 'r', encoding=CHASET) as f:
        return f.read()


# ===========================
# JS 配置更新函数
# ===========================
def update_js_config(sheets, section_sheets):
    """更新所有 JS 配置文件"""
    log.info("📝 更新 JS 配置文件...")
    
    # 从 index sheet 获取数据
    index_data = sheets.get('index', {}).get('data', [])
    
    # 更新各个 JS 文件
    update_seo_meta_js(index_data, section_sheets)
    update_count_badges_js(section_sheets)
    update_global_search_js(section_sheets)
    
    log.info("✓ JS 配置更新完成")


def update_seo_meta_js(index_data, section_sheets):
    """更新 seo-meta.js 中的影视系列列表"""
    seo_file = './docs/js/seo-meta.js'
    
    if not os.path.exists(seo_file):
        log.warn(f"⚠️ 找不到 {seo_file}，跳过")
        return
    
    with open(seo_file, 'r', encoding=CHASET) as f:
        content = f.read()
    
    # 生成影视系列列表 JavaScript
    series_list = []
    for idx, data in enumerate(index_data[:len(section_sheets)]):
        name = data.get('主标题', '')
        description = data.get('概要', '')
        if name:
            series_list.append(f"            {{ name: '{name}', description: '{description}' }}")
    
    series_js = ',\n'.join(series_list)
    
    # 替换 movieSeries 配置
    pattern = r"movieSeries:\s*\[[\s\S]*?\]"
    replacement = f"movieSeries: [\n{series_js}\n        ]"
    
    content = re.sub(pattern, replacement, content)
    
    with open(seo_file, 'w', encoding=CHASET) as f:
        f.write(content)
    
    log.info("✓ 已更新 docs/js/seo-meta.js")


def update_count_badges_js(section_sheets):
    """更新 count-badges.js 中的 SECTIONS 配置"""
    count_file = './docs/js/count-badges.js'
    
    if not os.path.exists(count_file):
        log.warn(f"⚠️ 找不到 {count_file}，跳过")
        return
    
    with open(count_file, 'r', encoding=CHASET) as f:
        content = f.read()
    
    # 生成 SECTIONS 配置
    sections_list = []
    for idx, sheet_name in enumerate(section_sheets):
        file_name = f'section-{idx+1:02d}.html'
        sections_list.append(f"        {{ file: '{file_name}', index: {idx} }}")
    
    sections_js = ',\n'.join(sections_list)
    
    # 替换 SECTIONS 配置
    pattern = r"const SECTIONS = \[[\s\S]*?\];"
    replacement = f"const SECTIONS = [\n{sections_js}\n    ];"
    
    content = re.sub(pattern, replacement, content)
    
    with open(count_file, 'w', encoding=CHASET) as f:
        f.write(content)
    
    log.info("✓ 已更新 docs/js/count-badges.js")


def update_global_search_js(section_sheets):
    """更新 global-search.js 中的 SECTIONS 配置"""
    search_file = './docs/js/global-search.js'
    
    if not os.path.exists(search_file):
        log.warn(f"⚠️ 找不到 {search_file}，跳过")
        return
    
    with open(search_file, 'r', encoding=CHASET) as f:
        content = f.read()
    
    # 生成 SECTIONS 配置（需要从 Excel 读取 section 名称）
    # 使用 sheet 名称作为 section 名称
    sections_list = []
    for idx, sheet_name in enumerate(section_sheets):
        file_name = f'section-{idx+1:02d}.html'
        sections_list.append(f"        {{ file: '{file_name}', name: '{sheet_name}' }}")
    
    sections_js = ',\n'.join(sections_list)
    
    # 替换 SECTIONS 配置
    pattern = r"const SECTIONS = \[[\s\S]*?\];"
    replacement = f"const SECTIONS = [\n{sections_js}\n    ];"
    
    content = re.sub(pattern, replacement, content)
    
    with open(search_file, 'w', encoding=CHASET) as f:
        f.write(content)
    
    log.info("✓ 已更新 docs/js/global-search.js")


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        log.error("❌ 发生未知异常")